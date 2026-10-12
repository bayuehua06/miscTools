#!/usr/bin/env python3
"""Markdown → APA 7th Word (.docx) CLI.

依赖: 仓库根目录运行 ./install.sh（或 pip install -r requirements.txt）

用法:
  ./misctools md2word paper.md
  ./misctools md2word paper.md -o paper.docx
  ./misctools md2word paper.md --author "Jane Smith" --title "My Study"

Markdown front matter（可选）:
  ---
  title: The Impact of Sleep on Memory
  author: Jane Smith
  affiliation: University of Example
  course: PSY 101: Introduction to Psychology
  instructor: Dr. John Doe
  date: October 12, 2026
  ---

APA 约定（学生论文）:
  - US Letter、四边 1 英寸边距
  - Times New Roman 12pt、全文双倍行距
  - 标题页（有元数据时）
  - 正文首行缩进 0.5"
  - Markdown #–##### → Word Heading 1–5（导航窗格可用），外观按 APA
  - 「References」节悬挂缩进
  - GFM 表格：表号（粗体）+ 斜体表题 + 仅顶线/表头下线/底线（无竖线）

表格 Markdown 示例:
  Table 1
  *Mean Digit-Span Scores by Sleep Group*

  | Group | M    | SD   |
  |-------|------|------|
  | Short | 5.21 | 1.10 |
  | Long  | 7.04 | 0.98 |

  *Note.* Scores are based on the digit-span task.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import List, Optional, Tuple

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt
except ImportError:
    print("错误: 未安装 python-docx。请在 miscTools 根目录执行: ./install.sh", file=sys.stderr)
    sys.exit(1)


FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(12)
MARGIN = Inches(1)
FIRST_LINE_INDENT = Inches(0.5)
HANGING_INDENT = Inches(0.5)


@dataclass
class Meta:
    title: str = ""
    author: str = ""
    affiliation: str = ""
    course: str = ""
    instructor: str = ""
    date: str = ""
    running_head: str = ""


@dataclass
class InlineSeg:
    text: str
    bold: bool = False
    italic: bool = False


@dataclass
class Block:
    kind: str  # heading|para|list_item|blockquote|blank|hr|table
    level: int = 0
    ordered: bool = False
    index: int = 0
    segs: List[InlineSeg] = field(default_factory=list)
    raw: str = ""
    # table fields
    rows: List[List[str]] = field(default_factory=list)
    table_number: str = ""
    table_title: str = ""
    table_note_segs: List[InlineSeg] = field(default_factory=list)


def is_pipe_row(line: str) -> bool:
    s = line.strip()
    return "|" in s and not re.match(r"^(-{3,}|\*{3,}|_{3,})\s*$", s)


def is_table_separator(line: str) -> bool:
    s = line.strip().strip("|").strip()
    if not s:
        return False
    parts = [p.strip() for p in s.split("|")]
    if not parts:
        return False
    return all(re.match(r"^:?-{3,}:?$", p) for p in parts if p)


def split_table_row(line: str) -> List[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def is_table_number_text(text: str) -> bool:
    t = text.strip()
    t = re.sub(r"^\*{1,2}|\*{1,2}$", "", t).strip()
    return bool(re.match(r"(?i)^table\s+\d+$", t)) or bool(
        re.match(r"^表\s*\d+$", t)
    )


def normalize_table_number(text: str) -> str:
    t = re.sub(r"^\*{1,2}|\*{1,2}$", "", text.strip()).strip()
    m = re.match(r"(?i)^table\s+(\d+)$", t)
    if m:
        return f"Table {m.group(1)}"
    m = re.match(r"^表\s*(\d+)$", t)
    if m:
        return f"Table {m.group(1)}"
    return t


def italic_only_text(segs: List[InlineSeg]) -> Optional[str]:
    if not segs:
        return None
    if all(s.italic and not s.bold for s in segs if s.text.strip()):
        return "".join(s.text for s in segs).strip()
    # Also accept whole line wrapped as *...*
    raw = "".join(s.text for s in segs).strip()
    m = re.match(r"^\*([^*]+)\*$", raw)
    if m:
        return m.group(1).strip()
    return None


def is_table_note_text(text: str) -> bool:
    t = text.strip()
    return bool(re.match(r"(?i)^\*?note\b", t)) or t.startswith("注") or t.startswith("说明")


def peek_table_caption(blocks: List[Block]) -> Tuple[str, str]:
    """Pull optional Table N / italic title from trailing blocks."""
    number = ""
    title = ""
    while blocks and blocks[-1].kind == "blank":
        blocks.pop()
    if blocks and blocks[-1].kind == "para":
        title_candidate = italic_only_text(blocks[-1].segs)
        raw = blocks[-1].raw.strip()
        if title_candidate:
            title = title_candidate
            blocks.pop()
            while blocks and blocks[-1].kind == "blank":
                blocks.pop()
        elif is_table_number_text(raw):
            number = normalize_table_number(raw)
            blocks.pop()
            return number, title
    if blocks and blocks[-1].kind == "para" and is_table_number_text(blocks[-1].raw):
        number = normalize_table_number(blocks[-1].raw)
        blocks.pop()
    return number, title


def parse_front_matter(text: str) -> Tuple[Meta, str]:
    meta = Meta()
    if not text.startswith("---"):
        return meta, text
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n?(.*)$", text, re.DOTALL)
    if not m:
        return meta, text
    body = m.group(2)
    for line in m.group(1).splitlines():
        line = line.strip()
        if not line or line.startswith("#") or ":" not in line:
            continue
        key, val = line.split(":", 1)
        key = key.strip().lower().replace("-", "_")
        val = val.strip().strip('"').strip("'")
        if hasattr(meta, key):
            setattr(meta, key, val)
    return meta, body


def parse_inlines(text: str) -> List[InlineSeg]:
    """Parse **bold**, *italic*, ***both***, `code`, [text](url)."""
    segs: List[InlineSeg] = []
    pattern = re.compile(
        r"(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
    )
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            segs.append(InlineSeg(text[pos : m.start()]))
        tok = m.group(0)
        if tok.startswith("***") and tok.endswith("***"):
            segs.append(InlineSeg(tok[3:-3], bold=True, italic=True))
        elif tok.startswith("**") and tok.endswith("**"):
            segs.append(InlineSeg(tok[2:-2], bold=True))
        elif tok.startswith("*") and tok.endswith("*"):
            segs.append(InlineSeg(tok[1:-1], italic=True))
        elif tok.startswith("`") and tok.endswith("`"):
            segs.append(InlineSeg(tok[1:-1]))
        elif tok.startswith("["):
            link_m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            if link_m:
                segs.append(InlineSeg(f"{link_m.group(1)} ({link_m.group(2)})"))
            else:
                segs.append(InlineSeg(tok))
        pos = m.end()
    if pos < len(text):
        segs.append(InlineSeg(text[pos:]))
    if not segs:
        segs.append(InlineSeg(text))
    return segs


def parse_markdown(body: str) -> List[Block]:
    blocks: List[Block] = []
    lines = body.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            blocks.append(Block(kind="blank"))
            i += 1
            continue

        if re.match(r"^(-{3,}|\*{3,}|_{3,})\s*$", stripped):
            blocks.append(Block(kind="hr"))
            i += 1
            continue

        hm = re.match(r"^(#{1,5})\s+(.+)$", stripped)
        if hm:
            level = len(hm.group(1))
            title = hm.group(2).strip()
            blocks.append(
                Block(kind="heading", level=level, segs=parse_inlines(title), raw=title)
            )
            i += 1
            continue

        um = re.match(r"^[-*+]\s+(.+)$", stripped)
        if um:
            blocks.append(
                Block(kind="list_item", ordered=False, segs=parse_inlines(um.group(1)))
            )
            i += 1
            continue

        om = re.match(r"^(\d+)[.)]\s+(.+)$", stripped)
        if om:
            blocks.append(
                Block(
                    kind="list_item",
                    ordered=True,
                    index=int(om.group(1)),
                    segs=parse_inlines(om.group(2)),
                )
            )
            i += 1
            continue

        if stripped.startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q = re.sub(r"^>\s?", "", lines[i].strip())
                quote_lines.append(q)
                i += 1
            text = " ".join(quote_lines)
            blocks.append(Block(kind="blockquote", segs=parse_inlines(text), raw=text))
            continue

        # GFM pipe table: header + separator + body rows
        if is_pipe_row(stripped) and i + 1 < len(lines) and is_table_separator(
            lines[i + 1]
        ):
            header = split_table_row(lines[i])
            i += 2
            rows = [header]
            while i < len(lines) and is_pipe_row(lines[i]) and not is_table_separator(
                lines[i]
            ):
                rows.append(split_table_row(lines[i]))
                i += 1
            # Normalize column count
            ncols = max((len(r) for r in rows), default=0)
            rows = [r + [""] * (ncols - len(r)) for r in rows]

            number, title = peek_table_caption(blocks)
            note_segs: List[InlineSeg] = []
            # Optional note immediately after table (skip one blank)
            j = i
            if j < len(lines) and not lines[j].strip():
                j += 1
            if j < len(lines) and is_table_note_text(lines[j].strip()):
                note_text = lines[j].strip()
                j += 1
                while j < len(lines) and lines[j].strip() and not is_pipe_row(
                    lines[j]
                ) and not re.match(r"^(#{1,5})\s+", lines[j].strip()):
                    # stop at structural; allow soft-wrapped note lines
                    if lines[j].strip().startswith(">") or re.match(
                        r"^[-*+]\s+", lines[j].strip()
                    ):
                        break
                    note_text += " " + lines[j].strip()
                    j += 1
                note_segs = parse_inlines(note_text)
                i = j

            blocks.append(
                Block(
                    kind="table",
                    rows=rows,
                    table_number=number,
                    table_title=title,
                    table_note_segs=note_segs,
                )
            )
            continue

        # Keep APA table number / italic title as their own blocks (even without blank lines)
        if is_table_number_text(stripped):
            blocks.append(
                Block(kind="para", segs=parse_inlines(stripped), raw=stripped)
            )
            i += 1
            continue
        if re.match(r"^\*[^*]+\*$", stripped) or re.match(r"^_[^_]+_$", stripped):
            blocks.append(
                Block(kind="para", segs=parse_inlines(stripped), raw=stripped)
            )
            i += 1
            continue

        para_parts = [stripped]
        i += 1
        while i < len(lines):
            ns = lines[i].strip()
            if not ns:
                break
            if re.match(r"^(#{1,5})\s+", ns):
                break
            if re.match(r"^[-*+]\s+", ns):
                break
            if re.match(r"^\d+[.)]\s+", ns):
                break
            if ns.startswith(">"):
                break
            if re.match(r"^(-{3,}|\*{3,}|_{3,})\s*$", ns):
                break
            if is_table_number_text(ns):
                break
            if re.match(r"^\*[^*]+\*$", ns) or re.match(r"^_[^_]+_$", ns):
                break
            # Do not swallow the start of a table into a paragraph
            if is_pipe_row(ns) and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
                break
            para_parts.append(ns)
            i += 1
        text = " ".join(para_parts)
        blocks.append(Block(kind="para", segs=parse_inlines(text), raw=text))

    # Auto-number tables that lack an explicit "Table N"
    auto_n = 1
    for b in blocks:
        if b.kind == "table" and not b.table_number:
            b.table_number = f"Table {auto_n}"
            auto_n += 1
        elif b.kind == "table" and b.table_number:
            m = re.search(r"(\d+)", b.table_number)
            if m:
                auto_n = max(auto_n, int(m.group(1)) + 1)

    return blocks


def _set_style_east_asia_font(style) -> None:
    style.font.name = FONT_NAME
    try:
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn("w:ascii"), FONT_NAME)
        rFonts.set(qn("w:hAnsi"), FONT_NAME)
        rFonts.set(qn("w:eastAsia"), FONT_NAME)
        rFonts.set(qn("w:cs"), FONT_NAME)
    except Exception:
        pass


def set_run_font(run, *, bold=False, italic=False, size=FONT_SIZE):
    run.font.name = FONT_NAME
    run.font.size = size
    run.bold = bold
    run.italic = italic
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), FONT_NAME)
    rFonts.set(qn("w:hAnsi"), FONT_NAME)
    rFonts.set(qn("w:eastAsia"), FONT_NAME)
    rFonts.set(qn("w:cs"), FONT_NAME)


def configure_apa_heading_styles(doc: Document) -> None:
    """Configure built-in Heading 1–5 for APA look + Word outline/navigation."""
    specs = {
        1: dict(bold=True, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, indent=Inches(0)),
        2: dict(bold=True, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, indent=Inches(0)),
        3: dict(bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, indent=Inches(0)),
        4: dict(bold=True, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, indent=FIRST_LINE_INDENT),
        5: dict(bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, indent=FIRST_LINE_INDENT),
    }
    for level, spec in specs.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT_NAME
        style.font.size = FONT_SIZE
        style.font.bold = spec["bold"]
        style.font.italic = spec["italic"]
        try:
            from docx.shared import RGBColor

            # Default Heading styles are often blue; APA uses black
            style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        except Exception:
            pass
        _set_style_east_asia_font(style)
        pf = style.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.alignment = spec["align"]
        pf.first_line_indent = spec["indent"]
        pf.left_indent = Inches(0)
        # Ensure outline level for Navigation pane (0-based)
        try:
            pPr = style.element.get_or_add_pPr()
            outline = pPr.find(qn("w:outlineLvl"))
            if outline is None:
                outline = OxmlElement("w:outlineLvl")
                pPr.append(outline)
            outline.set(qn("w:val"), str(level - 1))
        except Exception:
            pass


def set_paragraph_format(
    paragraph,
    *,
    align=None,
    first_line_indent=None,
    left_indent=None,
    space_before=Pt(0),
    space_after=Pt(0),
    hanging=False,
):
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = space_before
    pf.space_after = space_after
    if align is not None:
        paragraph.alignment = align
    if hanging:
        pf.left_indent = HANGING_INDENT
        pf.first_line_indent = -HANGING_INDENT
    else:
        if left_indent is not None:
            pf.left_indent = left_indent
        if first_line_indent is not None:
            pf.first_line_indent = first_line_indent


def add_page_number(paragraph):
    run = paragraph.add_run()
    set_run_font(run)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run2 = paragraph.add_run()
    set_run_font(run2)
    run2._r.append(instr)
    run3 = paragraph.add_run()
    set_run_font(run3)
    run3._r.append(fld_end)


def setup_document(doc: Document, meta: Meta):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN
    section.left_margin = MARGIN
    section.right_margin = MARGIN

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    set_paragraph_format(hp, align=WD_ALIGN_PARAGRAPH.RIGHT)
    if meta.running_head:
        # Professional papers: running head left + page right (simplified here).
        run = hp.add_run(meta.running_head.upper()[:50])
        set_run_font(run)
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        tab = hp.add_run("\t")
        set_run_font(tab)
    add_page_number(hp)

    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = FONT_SIZE
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    _set_style_east_asia_font(style)

    # Map Markdown #…##### → Word Heading 1–5（导航窗格 / 大纲可用），外观按 APA 7
    configure_apa_heading_styles(doc)


def add_inline_runs(paragraph, segs: List[InlineSeg]):
    for seg in segs:
        run = paragraph.add_run(seg.text)
        set_run_font(run, bold=seg.bold, italic=seg.italic)


def write_title_page(doc: Document, meta: Meta):
    for _ in range(3):
        p = doc.add_paragraph()
        set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER)

    title = meta.title or "Untitled Paper"
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = p.add_run(title)
    set_run_font(run, bold=True)

    blank = doc.add_paragraph()
    set_paragraph_format(blank, align=WD_ALIGN_PARAGRAPH.CENTER)

    for value in (
        meta.author,
        meta.affiliation,
        meta.course,
        meta.instructor,
        meta.date,
    ):
        if not value:
            continue
        p = doc.add_paragraph()
        set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(value)
        set_run_font(run)

    doc.add_page_break()


def is_references_heading(text: str) -> bool:
    t = text.strip().lower()
    return t in {"references", "reference", "参考文献", "参考资料", "works cited"}


def title_case_apa(text: str) -> str:
    if text != text.lower():
        return text
    small = {
        "a",
        "an",
        "the",
        "and",
        "but",
        "or",
        "for",
        "nor",
        "on",
        "at",
        "to",
        "from",
        "by",
        "of",
        "in",
        "with",
    }
    words = text.split()
    out = []
    for i, w in enumerate(words):
        if i == 0 or i == len(words) - 1 or w.lower() not in small:
            out.append(w[:1].upper() + w[1:] if w else w)
        else:
            out.append(w.lower())
    return " ".join(out)


def write_heading(doc: Document, block: Block):
    """Write Markdown heading as Word Heading N style (APA appearance)."""
    text = title_case_apa("".join(s.text for s in block.segs))
    level = max(1, min(int(block.level or 1), 5))

    if level >= 4 and not text.endswith("."):
        text += "."

    style_name = f"Heading {level}"
    p = doc.add_paragraph(style=style_name)
    # Re-assert APA paragraph format (some templates override)
    if level == 1:
        set_paragraph_format(
            p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=Inches(0)
        )
        run = p.add_run(text)
        set_run_font(run, bold=True)
    elif level == 2:
        set_paragraph_format(
            p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent=Inches(0)
        )
        run = p.add_run(text)
        set_run_font(run, bold=True)
    elif level == 3:
        set_paragraph_format(
            p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent=Inches(0)
        )
        run = p.add_run(text)
        set_run_font(run, bold=True, italic=True)
    elif level == 4:
        set_paragraph_format(
            p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent=FIRST_LINE_INDENT
        )
        run = p.add_run(text)
        set_run_font(run, bold=True)
    else:
        set_paragraph_format(
            p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent=FIRST_LINE_INDENT
        )
        run = p.add_run(text)
        set_run_font(run, bold=True, italic=True)


def set_cell_borders(cell, *, top=None, bottom=None, left=None, right=None):
    """Set cell borders. Values are (sz, color) where sz is eighths of a point, or None to clear."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is not None:
        tcPr.remove(tcBorders)
    tcBorders = OxmlElement("w:tcBorders")
    for edge, val in (
        ("top", top),
        ("left", left),
        ("bottom", bottom),
        ("right", right),
    ):
        el = OxmlElement(f"w:{edge}")
        if val is None:
            el.set(qn("w:val"), "nil")
        else:
            sz, color = val
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def apply_apa_table_borders(table):
    """APA: horizontal lines at top, under header, and bottom; no vertical lines."""
    n_rows = len(table.rows)
    n_cols = len(table.columns)
    # ½ pt ≈ sz 4; APA often uses 1–1½ pt; use sz=12 (~1.5pt) for clarity
    line = (12, "000000")
    for r in range(n_rows):
        for c in range(n_cols):
            cell = table.cell(r, c)
            top = line if r == 0 else None
            bottom = None
            if r == 0:
                bottom = line  # header bottom
            if r == n_rows - 1:
                bottom = line  # table bottom
            set_cell_borders(cell, top=top, bottom=bottom, left=None, right=None)


def clear_table_style_borders(table):
    """Remove default Word table grid so only our cell borders show."""
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)
    # Remove tblBorders if present
    for child in list(tblPr):
        if child.tag == qn("w:tblBorders"):
            tblPr.remove(child)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblPr.append(borders)


def write_table(doc: Document, block: Block):
    rows = block.rows
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    nrows = len(rows)

    # Table number (bold, flush left)
    if block.table_number:
        p = doc.add_paragraph()
        set_paragraph_format(
            p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent=Inches(0)
        )
        run = p.add_run(block.table_number)
        set_run_font(run, bold=True)

    # Table title (italic, flush left)
    if block.table_title:
        p = doc.add_paragraph()
        set_paragraph_format(
            p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent=Inches(0)
        )
        run = p.add_run(title_case_apa(block.table_title))
        set_run_font(run, italic=True)

    table = doc.add_table(rows=nrows, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    clear_table_style_borders(table)

    for r_idx, row_cells in enumerate(rows):
        for c_idx in range(ncols):
            text = row_cells[c_idx] if c_idx < len(row_cells) else ""
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            p = cell.paragraphs[0]
            # Tables: single spacing is common and readable; keep APA font
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.first_line_indent = Inches(0)
            segs = parse_inlines(text)
            if r_idx == 0:
                # Header row: bold
                for seg in segs:
                    run = p.add_run(seg.text)
                    set_run_font(run, bold=True, italic=seg.italic)
            else:
                add_inline_runs(p, segs)

    apply_apa_table_borders(table)

    # Note under table (flush left, no first-line indent; "Note." italic)
    if block.table_note_segs:
        p = doc.add_paragraph()
        set_paragraph_format(
            p, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_indent=Inches(0)
        )
        note_text = "".join(s.text for s in block.table_note_segs)
        m = re.match(r"(?i)^(\*?note\.?\*?\s*)(.*)$", note_text)
        if m:
            run = p.add_run("Note. ")
            set_run_font(run, italic=True)
            rest = m.group(2).lstrip()
            if rest:
                add_inline_runs(p, parse_inlines(rest))
        else:
            add_inline_runs(p, block.table_note_segs)


def write_body(
    doc: Document, blocks: List[Block], *, body_title: Optional[str] = None
):
    in_references = False
    started = False

    # APA 7：标题页之后，正文第一页顶端重复论文标题（居中加粗）
    if body_title:
        p = doc.add_paragraph(style="Heading 1")
        set_paragraph_format(
            p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=Inches(0)
        )
        run = p.add_run(body_title)
        set_run_font(run, bold=True)

    for block in blocks:
        if block.kind in {"blank", "hr"}:
            continue

        if block.kind == "heading":
            heading_text = "".join(s.text for s in block.segs).strip()
            if (
                not started
                and body_title
                and block.level == 1
                and heading_text.lower() == body_title.lower()
            ):
                continue
            started = True
            in_references = is_references_heading(heading_text)
            if in_references:
                p = doc.add_paragraph(style="Heading 1")
                set_paragraph_format(
                    p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=Inches(0)
                )
                run = p.add_run("References")
                set_run_font(run, bold=True)
            else:
                write_heading(doc, block)
            continue

        started = True

        if block.kind == "table":
            write_table(doc, block)
            continue

        if block.kind == "blockquote":
            p = doc.add_paragraph()
            set_paragraph_format(
                p,
                align=WD_ALIGN_PARAGRAPH.LEFT,
                left_indent=Inches(0.5),
                first_line_indent=Inches(0),
            )
            add_inline_runs(p, block.segs)
            continue

        if block.kind == "list_item":
            p = doc.add_paragraph()
            set_paragraph_format(
                p,
                align=WD_ALIGN_PARAGRAPH.LEFT,
                left_indent=Inches(0.5),
                first_line_indent=Inches(0),
            )
            prefix = f"{block.index}. " if block.ordered else "• "
            run = p.add_run(prefix)
            set_run_font(run)
            add_inline_runs(p, block.segs)
            continue

        if block.kind == "para":
            p = doc.add_paragraph()
            if in_references:
                set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.LEFT, hanging=True)
            else:
                set_paragraph_format(
                    p,
                    align=WD_ALIGN_PARAGRAPH.LEFT,
                    first_line_indent=FIRST_LINE_INDENT,
                )
            add_inline_runs(p, block.segs)


def has_title_page_meta(meta: Meta) -> bool:
    return bool(meta.title or meta.author or meta.affiliation or meta.course)


def merge_meta(file_meta: Meta, cli_meta: Meta, body: str, fallback_title: str) -> Meta:
    meta = Meta(**file_meta.__dict__)
    for key in (
        "title",
        "author",
        "affiliation",
        "course",
        "instructor",
        "date",
        "running_head",
    ):
        val = getattr(cli_meta, key)
        if val:
            setattr(meta, key, val)
    if not meta.title:
        for line in body.splitlines():
            hm = re.match(r"^#\s+(.+)$", line.strip())
            if hm:
                meta.title = hm.group(1).strip()
                break
    if not meta.title:
        meta.title = fallback_title
    return meta


def convert(md_path: Path, out_path: Path, cli_meta: Meta, *, no_title_page: bool) -> None:
    text = md_path.read_text(encoding="utf-8")
    file_meta, body = parse_front_matter(text)
    meta = merge_meta(
        file_meta,
        cli_meta,
        body,
        fallback_title=md_path.stem.replace("_", " ").replace("-", " "),
    )
    blocks = parse_markdown(body)

    doc = Document()
    setup_document(doc, meta)

    if no_title_page:
        write_body(doc, blocks)
    elif has_title_page_meta(meta):
        write_title_page(doc, meta)
        write_body(doc, blocks, body_title=meta.title)
    else:
        write_body(doc, blocks)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="md_to_apa_word",
        description="将 Markdown 转为符合 APA 7th（学生论文）版式的 Word 文档。",
    )
    p.add_argument("input", type=Path, help="输入 .md 文件")
    p.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="输出 .docx（默认与输入同目录同名）",
    )
    p.add_argument("--title", default="", help="论文标题（覆盖 front matter）")
    p.add_argument("--author", default="", help="作者")
    p.add_argument("--affiliation", default="", help="机构 / 学校")
    p.add_argument("--course", default="", help="课程名")
    p.add_argument("--instructor", default="", help="教师")
    p.add_argument("--date", default="", help="日期")
    p.add_argument(
        "--running-head",
        default="",
        dest="running_head",
        help="页眉 running head（专业稿可选）",
    )
    p.add_argument(
        "--no-title-page",
        action="store_true",
        help="不生成标题页（即使有元数据）",
    )
    return p


def main(argv: Optional[List[str]] = None) -> int:
    args = build_arg_parser().parse_args(argv)
    md_path = args.input.expanduser().resolve()
    if not md_path.is_file():
        print(f"错误: 找不到文件: {md_path}", file=sys.stderr)
        return 1
    if md_path.suffix.lower() not in {".md", ".markdown", ".txt"}:
        print(f"警告: 输入扩展名不是 .md: {md_path.suffix}", file=sys.stderr)

    out_path = (
        args.output.expanduser().resolve()
        if args.output
        else md_path.with_suffix(".docx")
    )
    cli_meta = Meta(
        title=args.title,
        author=args.author,
        affiliation=args.affiliation,
        course=args.course,
        instructor=args.instructor,
        date=args.date,
        running_head=args.running_head,
    )
    convert(md_path, out_path, cli_meta, no_title_page=args.no_title_page)
    print(f"完成: {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
